import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent / ".artifact_tools"))

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


SOURCE = Path(r"C:\Users\John Carlo Ocray\Downloads\[Reading Activity 2025] CTADMOBL Advanced Mobile Programming.docx")
OUTPUT = Path(__file__).parent / "Reading Activity ImagIngDev Answered.docx"

answers = [
    "The article talks about ImagIngDev, a tool that can turn a drawing of an app screen into app code. It can make code for different kinds of phones, so a beginner does not have to start each version from zero. The authors tested it and found that it was easy to learn and finished their sample app in fewer steps than another tool.",
    "I agree with the idea because drawing a screen first would help me plan what I want to make. In the article's test, ImagIng Tool took about four minutes while the other tool took about six, but the test did not include fixing the app later. I would still check the buttons and code because a drawing alone may not show everything the app needs.",
    "This connects to my Flutter class because I also make screens and think about how an app will look on different phones. A small shop could draw a simple product page and use a tool like this to make a first version of its app. I want to learn how to improve the code it makes so the app works well for real users.",
    "This reading made me see that planning an app with a drawing can make coding feel less hard. I still wonder if the tool can understand a messy drawing or a screen with many details. I would like to try it with one of my own designs and compare it with the same screen I build in Flutter.",
]


def set_cell_paragraph(cell, text):
    paragraph = cell.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.15
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(11)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


doc = Document(SOURCE)
table = doc.tables[0]

for row_index, answer in zip([1, 3, 5, 7], answers):
    set_cell_paragraph(table.cell(row_index, 0), answer)

# The user requested four answer paragraphs. Combine the final two question
# groups so their reflection and follow-up questions share the fourth answer.
final_prompt = table.cell(6, 0)
set_cell_paragraph(final_prompt, "Personal Reflection and Further Exploration")
for question in [
    "What did you learn about yourself or your own thinking from engaging with this reading?",
    "What questions do you still have after reading the article?",
    "What additional research would you like to do based on this reading?",
]:
    paragraph = final_prompt.add_paragraph(question)
    paragraph.paragraph_format.space_after = Pt(2)

for row_index in [9, 8]:
    row_element = table.rows[row_index]._element
    row_element.getparent().remove(row_element)

doc.save(OUTPUT)
print(OUTPUT)
